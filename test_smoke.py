"""Isolated smoke checks: no real receipts, credentials or provider calls."""
import json
from pathlib import Path
import tempfile
import threading
import unittest
import urllib.request
import urllib.error
from http.server import ThreadingHTTPServer
from unittest.mock import patch

import serve
from export_expense import export
from openpyxl import load_workbook


class SmokeTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.db = self.root / 'data.db'
        self.patches = [patch.object(serve, 'DB_PATH', self.db),
                        patch.object(serve, 'DATA_DIR', self.root),
                        patch.object(serve, 'INVOICES_DIR', self.root / 'invoices'),
                        patch.object(serve, 'EXPORTS_DIR', self.root / 'exports')]
        for p in self.patches:
            p.start()
        serve.init_db()

    def tearDown(self):
        for p in reversed(self.patches):
            p.stop()
        self.tmp.cleanup()

    def test_empty_database_and_generic_identity(self):
        db = serve.get_db()
        self.assertEqual(db.execute('SELECT COUNT(*) FROM records').fetchone()[0], 0)
        self.assertEqual(db.execute("SELECT value FROM settings WHERE key='user_name'").fetchone()[0], '报销用户')
        self.assertEqual(db.execute("SELECT value FROM settings WHERE key='api_key'").fetchone()[0], '')
        db.close()

    def test_export_synthetic_record(self):
        records = [{'发生日期': '2026-10-01', '招待客户名称': '示例客户', '人数': 2, '金额': 128.5}]
        result = export('10', records_override=records, return_result=True,
                        exports_dir=self.root / 'exports', db_path=self.db)
        self.assertTrue(result['ok'], result)
        wb = load_workbook(result['path'])
        ws = wb['费用报销清单']
        self.assertEqual(ws['D3'].value, '示例客户')
        self.assertEqual(ws['I3'].value, 128.5)
        self.assertEqual(ws['I15'].value, '=SUM(I3:I3)')
        wb.close()

    def test_database_and_source_not_served(self):
        server = ThreadingHTTPServer(('127.0.0.1', 0), serve.Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        base = f'http://127.0.0.1:{server.server_port}'
        try:
            with urllib.request.urlopen(base + '/api/config', timeout=3) as response:
                self.assertEqual(response.status, 200)
                json.load(response)
            for path in ['/data.db', '/feishu_config.json', '/serve.py', '/.git/config']:
                with self.assertRaises(urllib.error.HTTPError) as error:
                    urllib.request.urlopen(base + path, timeout=3)
                self.assertEqual(error.exception.code, 404)
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=3)


if __name__ == '__main__':
    unittest.main()
